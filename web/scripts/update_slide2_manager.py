from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

path = r"d:\Project\Digitalisasi Compound Slow Moving\Alur Compound Slow Moving.pptx"
prs = Presentation(path)
slide = prs.slides[1]

INK = RGBColor(0x1A, 0x17, 0x14)
RED = RGBColor(0xC3, 0x1F, 0x16)
OK = RGBColor(0x1F, 0x5C, 0x3C)
MUTED = RGBColor(0x4A, 0x45, 0x3C)
AMBER = RGBColor(0x8A, 0x55, 0x12)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GOLD = RGBColor(0xE8, 0xC9, 0x8A)


def set_lines(shape, lines, size=10, bold=False, color=INK, align=PP_ALIGN.CENTER, title=None, title_size=13, title_color=None):
    tf = shape.text_frame
    tf.word_wrap = True
    items = []
    if title:
        items.append((title, title_size, True, title_color or color, align if align != PP_ALIGN.CENTER else PP_ALIGN.LEFT))
    for line in lines:
        items.append((line, size, bold, color, align))

    # reuse existing paragraphs, drop extras
    paras = list(tf.paragraphs)
    while len(paras) > 1:
        p = paras.pop()
        p._p.getparent().remove(p._p)

    first = True
    for text, sz, is_bold, col, al in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = al
        p.space_after = Pt(1)
        p.space_before = Pt(0)
        # clear runs
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        run = p.add_run()
        run.text = text
        run.font.size = Pt(sz)
        run.font.bold = is_bold
        run.font.color.rgb = col
        run.font.name = "Aptos"


s = slide.shapes

set_lines(
    s[1],
    ["Alur to-be  |  foto tiket + konfirmasi, 1 label saat disposisi, eksekusi kg, progress maju/mundur"],
    size=12,
    color=MUTED,
    align=PP_ALIGN.LEFT,
)

set_lines(
    s[2],
    ["Sumber SM + leftover/New Creation hari ini (jalur cepat ke langkah 1):  BB1, BB2, BB3, BB4, 10in Ext, 8in Ext, Triplex, 3 Roll Cal, 4 Roll Cal, UB Bead, Hexa Bead"],
    size=11,
    align=PP_ALIGN.CENTER,
)

set_lines(
    s[4],
    [
        "Foto Mixer Ticket di HP",
        "ML baca field tiket",
        "Operator konfirmasi hasil",
        "Timbang netto (wajib)",
        "Tiket tidak memuat kg",
        "Status: Under Review",
        "PIC: Mixing",
    ],
    size=10,
)

set_lines(s[5], ["Ganti isi Excel NEW"], size=9, bold=True, color=AMBER)

set_lines(
    s[8],
    [
        "Stok live: #skid & kg",
        "ID: compound + run + skid",
        "Expiry / retest terlihat",
        "QC cek kondisi material",
        "Mixing jaga lokasi rak",
        "Tidak nunggu recap malam",
    ],
    size=10,
)

set_lines(s[9], ["Ganti cek fisik + Excel"], size=9, bold=True, color=AMBER)

set_lines(
    s[12],
    [
        "Restacking / Pickle Out",
        "Pilih lokasi & kondisi",
        "Submit di HP saja",
        "Belum print, belum ttd",
        "Selesai → antri disposisi",
        "PIC: Mixing (fisik)",
    ],
    size=10,
)

set_lines(s[13], ["Tanpa label — update status"], size=9, bold=True, color=AMBER)

set_lines(
    s[16],
    [
        "QC: layak / kondisi",
        "PC: aturan pakai & jadwal",
        "Ttd = login NIK",
        "(bukan pilih nama orang)",
        "Print barcode + tempel",
        "Verify: SCAN label tempel",
        "Foto hanya untuk audit",
        "Scrap: Supervisor",
    ],
    size=9,
)

set_lines(s[17], ["Satu-satunya titik print"], size=9, bold=True, color=AMBER)

set_lines(
    s[20],
    [
        "Scan label di Ext/Cal/Mix",
        "Catat kg aktual terpakai",
        "Sisa skid tetap hidup",
        "ODS 10–20 kg/batch OK",
        "Gagal trial → Rak Merah",
        "PIC: area mesin",
    ],
    size=10,
)

set_lines(s[21], ["Tutup loop kg & sisa"], size=9, bold=True, color=AMBER)

set_lines(
    s[22],
    ["Dipakai ulang ke:  Triplex, 10in Ext, 8in Ext, 3 Roll Cal, 4 Roll Cal, UB Bead, Hexa Bead     |     Cabang: Cleaner / Scrap / kembali Rak Merah     |     Label fisik hanya terbit di langkah 4"],
    size=11,
    align=PP_ALIGN.CENTER,
)

set_lines(
    s[23],
    ["6. PAPAN PROGRESS HARIAN  —  satu stok, PIC jelas, leftover terdaftar hari itu juga"],
    size=12,
    bold=True,
    color=WHITE,
    align=PP_ALIGN.LEFT,
)

set_lines(
    s[24],
    [
        "Total kg & skid Rak Merah turun vs kemarin",
        "Plan harian PC tercapai (blend / ODS / pakai)",
        "Tag tereksekusi sebelum expired",
        "New Creation masuk sistem di hari yang sama",
        "Label tertempel dan sudah di-scan verify",
    ],
    size=11,
    align=PP_ALIGN.LEFT,
    title="MAJU  (progress positif)",
    title_color=OK,
)

set_lines(
    s[25],
    [
        "Leftover naik / mengendap karena tidak didaftar",
        "Eksekusi gagal, kg tidak terpotong, tag OPEN",
        "Ttd memakai nama orang lain / tanpa login",
        "Print dobel di restacking dan disposisi",
        "Fisik rak tidak sama dengan sistem",
    ],
    size=11,
    align=PP_ALIGN.LEFT,
    title="MUNDUR  (progress negatif)",
    title_color=RED,
)

prs.save(path)
print("updated slide 2 text")
